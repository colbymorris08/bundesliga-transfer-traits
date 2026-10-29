#!/usr/bin/env python3
"""Write the downloadable Word document of analysis sources."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "asian-pitcher-sources.docx"
NAVY = RGBColor(0x3D, 0x5A, 0x80)
PAGE = "https://colbymorris08.github.io/bundesliga-transfer-traits/asian-pitchers.html"


def add_link(paragraph, url):
    """Append a clickable URL whose visible text is the full address."""
    part = paragraph.part
    rel = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = paragraph._p.makeelement(qn("w:hyperlink"))
    hyperlink.set(qn("r:id"), rel)
    run = paragraph._p.makeelement(qn("w:r"))
    rpr = paragraph._p.makeelement(qn("w:rPr"))
    color = rpr.makeelement(qn("w:color"))
    color.set(qn("w:val"), "3D5A80")
    underline = rpr.makeelement(qn("w:u"))
    underline.set(qn("w:val"), "single")
    sz = rpr.makeelement(qn("w:sz"))
    sz.set(qn("w:val"), "22")
    rpr.append(color)
    rpr.append(underline)
    rpr.append(sz)
    run.append(rpr)
    text = run.makeelement(qn("w:t"))
    text.set(qn("xml:space"), "preserve")
    text.text = url
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def body(paragraph):
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def heading(doc, text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def bullet(doc, parts):
    """parts is a list of strings and ('link', url) tuples, in order."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    for part in parts:
        if isinstance(part, tuple) and part[0] == "link":
            add_link(p, part[1])
        else:
            run = p.add_run(part)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    doc.core_properties.title = "Sources for the Asian League to MLB Pitcher Analysis"
    doc.core_properties.subject = "Roster, rate, API, ERA+, bWAR, wOBA, and wRC+ sources"
    doc.core_properties.category = "References"

    title = doc.add_heading("Sources for the Asian League to MLB Pitcher Analysis", level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1C, 0x1C, 0x1C)
        run.font.name = "Calibri"

    intro = doc.add_paragraph()
    intro.add_run("October 29, 2026. Public page: ")
    add_link(intro, PAGE)
    body(intro)

    note = doc.add_paragraph()
    run = note.add_run(
        "Sources for the roster, the rates, and the calculations on that page. "
        "wOBA and wRC+ are listed after that. This page does not compute them. "
        "The run-prevention outcomes here are ERA+ and bWAR."
    )
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.italic = True
    body(note)

    heading(doc, "Roster and Asian-league seasons")
    bullet(doc, [
        "Baseball-Reference. Register player pages, standard pitching table. ",
        ("link", "https://www.baseball-reference.com/register/"),
        ". Pre-move NPB (Central and Pacific), KBO, and CPBL seasons. Farm and winter leagues are excluded. A season in the departure year counts as pre-move. A later return to Asia does not.",
    ])
    bullet(doc, [
        "Baseball-Reference. Birthplace pitching tables, used to check the Asian-born half of the roster. Japan: ",
        ("link", "https://www.baseball-reference.com/bio/Japan_born.shtml"),
        ". South Korea: ",
        ("link", "https://www.baseball-reference.com/bio/South-Korea_born.shtml"),
        ". Taiwan: ",
        ("link", "https://www.baseball-reference.com/bio/Taiwan_born.shtml"),
        ".",
    ])
    bullet(doc, [
        "Published window: the move, or the return to MLB, is in 2006 or later. A pitcher is in the correlations with a pre-move Asian season, at least 10 Asian innings before the move, and at least 10 MLB innings.",
    ])

    heading(doc, "MLB pitching stats")
    bullet(doc, [
        "MLB Advanced Media. MLB Stats API, career pitching. ",
        ("link", "https://statsapi.mlb.com/api/v1/people/660271/stats?stats=career&group=pitching"),
        ". stats=career and group=pitching.",
    ])
    bullet(doc, [
        "MLB Advanced Media. MLB Stats API, year-by-year pitching. stats=yearByYear and group=pitching, sport id 1, seasons from 1960 on. A blank-team season total is dropped when named team splits exist, so innings are not double-counted.",
    ])
    bullet(doc, [
        "Rates from those counting stats: K/9 = SO × 9 / IP, BB/9 = BB × 9 / IP, HR/9 = HR × 9 / IP, WHIP = (H + BB) / IP, K/BB = SO / BB. IP is the baseball decimal as published (242.1), not converted to true thirds.",
    ])

    heading(doc, "Ground-ball rate")
    bullet(doc, [
        "Baseball Savant. Batted-ball leaderboard, pitcher, 2015–2026. ",
        ("link", "https://baseballsavant.mlb.com/leaderboard/batted-ball?type=pitcher"),
        ". MLB GB% = ground balls / batted balls. A career rate is weighted by batted balls. The Stats API has ground outs and air outs only, and that ratio is not used.",
    ])
    bullet(doc, [
        "NPB Data Store. Pitch-type pages, ゴロ率 labeled GB%. ",
        ("link", "https://baseball-data-store.com/"),
        ". Used only when the page title is the requested season. The site currently publishes 2025 and 2026 for players it still lists.",
    ])
    bullet(doc, [
        "NPB official pitching lines do not publish ground-ball rate. Example, Central League pitching, 2026: ",
        ("link", "https://npb.jp/bis/2026/stats/pit_c.html"),
        ". KBO and CPBL public lines are the same, so those GB% cells are blank.",
    ])

    heading(doc, "ERA+, WAR, and the correlation")
    bullet(doc, [
        "Baseball-Reference. ERA+ (adjusted ERA). ",
        ("link", "https://www.baseball-reference.com/bullpen/ERA+"),
        ". Park- and league-adjusted ERA, scaled so 100 is average. Values on this page are the pitching bWAR file’s ERA+, averaged across MLB seasons.",
    ])
    bullet(doc, [
        "Baseball-Reference. Pitching WAR. ",
        ("link", "https://www.baseball-reference.com/about/war_explained_pitch.shtml"),
        ". Career WAR is the sum of seasonal bWAR from the Baseball-Reference pitching file, read with pybaseball bwar_pitch().",
    ])
    bullet(doc, [
        "Pearson, Karl. 1895. “Note on Regression and Inheritance in the Case of Two Parents.” Proceedings of the Royal Society of London 58: 240–242. Computed with scipy.stats.pearsonr. A metric is marked stable when |r| ≥ 0.50 and p < 0.10.",
    ])
    bullet(doc, [
        "Pedregosa, F., et al. 2011. “Scikit-learn: Machine Learning in Python.” Journal of Machine Learning Research 12: 2825–2830. Redundancy clusters use average-linkage agglomerative clustering on distance 1 − |r|, distance threshold 0.40. A pair is redundant at |r| ≥ 0.75. GB% is left out of that matrix.",
    ])
    bullet(doc, [
        "Cho, Andrew. 2026. “Cross-League Equivalencies: Using Former MLB Player Data to Evaluate NPB Transfers.” ",
        ("link", "https://acho25.me/QSS82_AndrewCho.pdf"),
        ". The 1.096 xFIP− factor is cited here and is not the stability result.",
    ])

    heading(doc, "wOBA and wRC+")
    caveat = doc.add_paragraph()
    run = caveat.add_run(
        "Published formulas. This analysis does not compute wOBA or wRC+. The outcomes on the page are ERA+ and bWAR."
    )
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    body(caveat)

    bullet(doc, [
        "Tango, Tom M., Mitchel G. Lichtman, and Andrew E. Dolphin. 2007. The Book: Playing the Percentages in Baseball. Washington, DC: Potomac Books. Linear weights behind wOBA.",
    ])
    bullet(doc, [
        "FanGraphs. “wOBA.” ",
        ("link", "https://library.fangraphs.com/offense/woba/"),
        ". wOBA = (wBB·uBB + wHBP·HBP + w1B·1B + w2B·2B + w3B·3B + wHR·HR) / (AB + BB − IBB + SF + HBP). Weights change by season.",
    ])
    bullet(doc, [
        "FanGraphs. Guts! constants, the season weights and the wOBA scale. ",
        ("link", "https://www.fangraphs.com/tools/guts"),
        ".",
    ])
    bullet(doc, [
        "FanGraphs. “wRC and wRC+.” ",
        ("link", "https://library.fangraphs.com/offense/wrc/"),
        ". wRC = [((wOBA − league wOBA) / wOBA scale) + (league R/PA)] × PA. wRC+ scales park-adjusted wRC per plate appearance to 100 for a league-average hitter.",
    ])

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
